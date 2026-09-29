"""Build the donations Excel workbook."""
from __future__ import annotations

import os
import re
from dataclasses import dataclass
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

from .links import gmail_url
from .models import NEEDS_REVIEW, Donation, Email

HEADERS = ["Date", "Organization", "Country", "Currency", "Amount", "Role", "Receipt #",
           "Charity ID", "Account", "Status", "Gmail", "Email file", "Receipt file", "Notes"]
YELLOW = PatternFill("solid", fgColor="FFFF00")
BOLD = Font(bold=True)
MONEY = "#,##0.00"


@dataclass
class ReportRow:
    date: date
    organization: str
    country: str
    currency: str
    amount: Decimal | None
    role: str
    receipt_no: str | None
    charity_id: str | None
    account: str
    status: str
    gmail_url: str
    email_file: Path | None
    receipt_file: Path | None
    notes: str


@dataclass
class Summary:
    totals: dict[tuple[str, str], Decimal]
    by_org: dict[tuple[str, str], Decimal]
    needs_review: int
    needs_rule: int
    skipped: int
    confirmation_only: list[str]  # organizations with confirmation rows but no counted receipt


def row_from_donation(
    d: Donation,
    email_file: Path | None = None,
    receipt_files: list[Path] | None = None,
    extra_notes: list[str] | None = None,
) -> ReportRow:
    receipts = d.receipt_paths if receipt_files is None else receipt_files
    notes = list(d.notes) + list(extra_notes or [])
    if len(receipts) > 1:
        notes.append("Other receipt files: " + ", ".join(p.name for p in receipts[1:]))
    return ReportRow(
        date=d.email.date.date(),
        organization=d.rule.name,
        country=d.rule.country,
        currency=d.rule.currency,
        amount=d.amount,
        role=d.rule.role,
        receipt_no=d.receipt_no,
        charity_id=d.charity_id,
        account=d.email.account,
        status=d.status,
        gmail_url=d.gmail_url,
        email_file=email_file or d.email_html_path,
        receipt_file=receipts[0] if receipts else None,
        notes="; ".join(notes),
    )


def summarize(donations: list[Donation], needs_rule: int, skipped: int) -> Summary:
    totals: dict[tuple[str, str], Decimal] = {}
    by_org: dict[tuple[str, str], Decimal] = {}
    for d in donations:
        if not d.is_countable:
            continue
        key = (d.rule.country, d.rule.currency)
        totals[key] = totals.get(key, Decimal(0)) + d.amount
        org_key = (d.rule.name, d.rule.currency)
        by_org[org_key] = by_org.get(org_key, Decimal(0)) + d.amount
    needs_review = sum(1 for d in donations if d.status == NEEDS_REVIEW)
    counted = {_organization(d.rule.name) for d in donations if d.is_countable}
    confirmed = {_organization(d.rule.name) for d in donations if d.rule.role == "confirmation"}
    return Summary(totals, by_org, needs_review, needs_rule, skipped, sorted(confirmed - counted))


def _organization(rule_name: str) -> str:
    """'ICNA Relief Canada (card payment)' -> 'ICNA Relief Canada', pairing a confirmation rule with its receipt rule."""
    return re.sub(r"\s*\([^)]*\)\s*$", "", rule_name)


def _link(cell, target: str, label: str) -> None:
    cell.value = label
    cell.hyperlink = target
    cell.style = "Hyperlink"


def _file_link(target: Path, workbook_path: Path, absolute: bool) -> str:
    # Absolute file:/// URIs work in viewers that don't resolve relative links (e.g. VS Code Office Viewer);
    # relative links keep a folder portable when zipped (accountant pack).
    if absolute:
        return target.resolve().as_uri()
    return os.path.relpath(target, workbook_path.parent).replace("\\", "/")


def _write_summary(ws, summary: Summary) -> None:
    ws.append(["Totals (receipts with status OK)"])
    ws["A1"].font = BOLD
    ws.append(["Country / Currency", "Total"])
    for (country, currency), total in sorted(summary.totals.items()):
        ws.append([f"{country} / {currency}", total])
    ws.append([])
    ws.append(["By organization"])
    ws.cell(ws.max_row, 1).font = BOLD
    ws.append(["Organization", "Currency", "Total"])
    for (org, currency), total in sorted(summary.by_org.items()):
        ws.append([org, currency, total])
    ws.append([])
    ws.append(["Checks"])
    ws.cell(ws.max_row, 1).font = BOLD
    ws.append(["Rows needing review", summary.needs_review])
    ws.append(["Emails needing a rule", summary.needs_rule])
    ws.append(["Malformed messages skipped", summary.skipped])
    ws.append(["Orgs with confirmations but no counted receipt", ", ".join(summary.confirmation_only) or "(none)"])
    for row in ws.iter_rows():
        for cell in row:
            if isinstance(cell.value, Decimal):
                cell.number_format = MONEY
    ws.column_dimensions["A"].width = 40


def _write_donations(ws, rows: list[ReportRow], workbook_path: Path, absolute_links: bool) -> None:
    ws.append(HEADERS)
    for cell in ws[1]:
        cell.font = BOLD
    for r in rows:
        ws.append([r.date, r.organization, r.country, r.currency, r.amount, r.role, r.receipt_no,
                   r.charity_id, r.account, r.status, None, None, None, r.notes])
        excel_row = ws[ws.max_row]
        excel_row[0].number_format = "yyyy-mm-dd"
        excel_row[4].number_format = MONEY
        if r.gmail_url:
            _link(excel_row[10], r.gmail_url, "Open in Gmail")
        if r.email_file:
            _link(excel_row[11], _file_link(r.email_file, workbook_path, absolute_links), r.email_file.name)
        if r.receipt_file:
            _link(excel_row[12], _file_link(r.receipt_file, workbook_path, absolute_links), r.receipt_file.name)
        if r.status == NEEDS_REVIEW:
            for cell in excel_row:
                cell.fill = YELLOW
    widths = [12, 30, 8, 9, 12, 13, 18, 20, 26, 14, 14, 30, 30, 50]
    for index, width in enumerate(widths):
        ws.column_dimensions[chr(ord("A") + index)].width = width
    ws.freeze_panes = "A2"


def _write_needs_rule(ws, emails: list[Email], year: int) -> None:
    ws.append(["Date", "Account", "From", "Subject", "Gmail", "Note"])
    for cell in ws[1]:
        cell.font = BOLD
    for email in sorted(emails, key=lambda e: e.date):
        note = f"Dated {year + 1}; possibly a receipt for {year}" if email.date.year == year + 1 else ""
        ws.append([email.date.date(), email.account, email.sender, email.subject, None, note])
        url = gmail_url(email.account, email.message_id)
        if url:
            _link(ws.cell(ws.max_row, 5), url, "Open in Gmail")
    for column, width in zip("ABCDEF", [12, 26, 45, 60, 14, 40]):
        ws.column_dimensions[column].width = width


def write_workbook(path: Path, rows: list[ReportRow], summary: Summary, needs_rule: list[Email] | None,
                   year: int, absolute_links: bool = False) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _write_summary(wb.active, summary)
    wb.active.title = "Summary"
    _write_donations(wb.create_sheet("Donations"), rows, path, absolute_links)
    if needs_rule is not None:
        _write_needs_rule(wb.create_sheet("Needs a rule"), needs_rule, year)
    wb.save(path)
