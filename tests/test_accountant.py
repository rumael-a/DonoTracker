from decimal import Decimal

from openpyxl import load_workbook

from donation_tracker.accountant import build_pack
from donation_tracker.models import NEEDS_REVIEW, OK, Donation
from donation_tracker.report import summarize
from tests.helpers import email_record, rule_record


class FakeRenderer:
    def __init__(self, ok=True):
        self.ok = ok

    def render(self, html_path, pdf_path):
        if self.ok:
            pdf_path.parent.mkdir(parents=True, exist_ok=True)
            pdf_path.write_bytes(b"%PDF-fake")
        return self.ok


def saved_donation(tmp_path, name, amount, *, role="receipt", status=OK, country="CA", currency="CAD", receipts=1):
    folder = tmp_path / "evidence" / name
    folder.mkdir(parents=True)
    html = folder / "email.html"
    html.write_text("<p>x</p>", encoding="utf-8")
    d = Donation(email_record(), rule_record(name=name, role=role, country=country, currency=currency),
                 Decimal(amount) if amount else None, None, None, status, [])
    d.email_html_path = html
    for i in range(receipts):
        pdf = folder / f"r{i}.pdf"
        pdf.write_bytes(b"%PDF")
        d.receipt_paths.append(pdf)
    return d


def test_pack_contains_only_ok_receipts_split_by_country(tmp_path):
    ds = [
        saved_donation(tmp_path, "Muslim Hands", "1200", receipts=2),
        saved_donation(tmp_path, "Red Cross", "250", country="US", currency="USD", receipts=0),
        saved_donation(tmp_path, "MyTenNights nightly", "3.75", role="confirmation"),
        saved_donation(tmp_path, "Broken", None, status=NEEDS_REVIEW),
    ]
    pack = tmp_path / "for_accountant"
    log = []
    files = build_pack(ds, pack, FakeRenderer(), 2025, summarize(ds, 0, 0), log)
    names = sorted(p.relative_to(pack).as_posix() for p in files)
    assert names == [
        "CA/2025-03-14_Muslim_Hands_CAD_1200.00_email.pdf",
        "CA/2025-03-14_Muslim_Hands_CAD_1200.00_receipt.pdf",
        "CA/2025-03-14_Muslim_Hands_CAD_1200.00_receipt_2.pdf",
        "US/2025-03-14_Red_Cross_USD_250.00_email.pdf",
    ]
    wb = load_workbook(pack / "donations_2025_summary.xlsx")
    assert wb.sheetnames == ["Summary", "Donations"]
    ws = wb["Donations"]
    assert ws.max_row == 3
    assert ws["L2"].hyperlink.target == "CA/2025-03-14_Muslim_Hands_CAD_1200.00_email.pdf"
    assert ws["M2"].hyperlink.target == "CA/2025-03-14_Muslim_Hands_CAD_1200.00_receipt.pdf"


def test_render_failure_falls_back_to_html(tmp_path):
    d = saved_donation(tmp_path, "Muslim Hands", "10", receipts=0)
    log = []
    files = build_pack([d], tmp_path / "pack", FakeRenderer(ok=False), 2025, summarize([d], 0, 0), log)
    assert [p.name for p in files] == ["2025-03-14_Muslim_Hands_CAD_10.00_email.html"]
    assert any("PDF render failed" in line for line in log)
    ws = load_workbook(tmp_path / "pack" / "donations_2025_summary.xlsx")["Donations"]
    assert "PDF render failed" in ws["N2"].value
