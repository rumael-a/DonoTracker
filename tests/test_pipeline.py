from datetime import datetime, timezone
from decimal import Decimal
from pathlib import Path
from urllib.parse import urlparse
from urllib.request import url2pathname

import pytest
from openpyxl import load_workbook

from donation_tracker.config import Account
from donation_tracker.detector import load_detector
from donation_tracker.models import NEEDS_REVIEW, OK, Donation
from donation_tracker.pipeline import flag_possible_duplicates, run
from donation_tracker.report import summarize
from tests.helpers import email_record, make_email, make_mbox, make_text_pdf, rule_record
from tests.test_accountant import FakeRenderer

ROOT = Path(__file__).resolve().parents[1]


def build_accounts(tmp_path):
    shared = make_email(sender="Muslim Hands <mail@muslimhands.ca>", subject="Your Receipt from Muslim Hands",
                        html="<p>Donation Amount: $1,200.00</p><p>Receipt No: MH-778</p>",
                        message_id="<mh1@muslimhands.ca>", date="Fri, 17 Oct 2025 10:00:00 -0400")
    a = [
        shared,
        make_email(sender="MyTenNights <contact@mytennights.com>", subject="Your 21st Night Donation",
                   text="Tonight's donation $7.17", message_id="<n21@mtn>", date="Fri, 21 Mar 2025 22:00:00 +0000"),
        make_email(sender="MyTenNights <contact@mytennights.com>", subject="Your Annual Tax Receipt",
                   text="See attached", message_id="<annual@mtn>", date="Tue, 24 Feb 2026 09:00:00 +0000",
                   attachments=[("MTN2025.pdf", "application/pdf", make_text_pdf("Total donated $54.60"))]),
        make_email(sender="MyTenNights <contact@mytennights.com>", subject="Your Annual Tax Receipt",
                   text="Total donated $99.00", message_id="<annual-prev@mtn>", date="Mon, 24 Feb 2025 09:00:00 +0000"),
        make_email(sender="Human Appeal USA <customercare@humanappeal.org>", subject="Thank you",
                   text="No amount here", message_id="<ha@x>", date="Sat, 21 Jun 2025 09:00:00 +0000"),
        make_email(sender="New Charity <hello@newcharity.org>", subject="Thank you for your donation",
                   text="We received it", message_id="<new@x>", date="Mon, 05 May 2025 09:00:00 +0000"),
        make_email(sender="Friend <f@gmail.com>", subject="Lunch", text="hi", message_id="<f@x>",
                   date="Mon, 05 May 2025 09:00:00 +0000"),
    ]
    b = [shared]  # same Message-ID in second mailbox → de-duplicated
    return [
        Account("abbasirumael@gmail.com", make_mbox(tmp_path / "a.mbox", a)),
        Account("itsrumzy@gmail.com", make_mbox(tmp_path / "b.mbox", b)),
    ]


def build_rules():
    return [
        rule_record(name="Muslim Hands Canada", match_from=["muslimhands.ca"], match_subject=["Your Receipt"]),
        rule_record(name="MyTenNights nightly", country="US", currency="USD", role="confirmation",
                    match_from=["mytennights.com"], match_subject=["Night Donation"]),
        rule_record(name="MyTenNights annual", country="US", currency="USD", tax_year="previous",
                    match_from=["mytennights.com"], match_subject=["Annual Tax Receipt"]),
        rule_record(name="Human Appeal USA", country="US", currency="USD", match_from=["humanappeal"]),
    ]


def test_end_to_end(tmp_path):
    out_root = tmp_path / "output"
    (out_root / "2025").mkdir(parents=True)
    (out_root / "2025" / "stale.txt").write_text("old", encoding="utf-8")
    result = run(2025, build_accounts(tmp_path), build_rules(), load_detector(ROOT / "detector.yaml"),
                 out_root, FakeRenderer())

    assert not (out_root / "2025" / "stale.txt").exists()
    names = [d.rule.name for d in result.donations]
    assert names == ["MyTenNights nightly", "Human Appeal USA", "Muslim Hands Canada", "MyTenNights annual"]
    annual = result.donations[-1]
    assert annual.amount == Decimal("54.60")
    assert annual.email.date.year == 2026
    assert result.summary.totals == {("CA", "CAD"): Decimal("1200.00"), ("US", "USD"): Decimal("54.60")}
    assert result.summary.needs_review == 1
    assert [e.message_id for e in result.needs_rule] == ["<new@x>"]
    assert result.donations[2].gmail_url.endswith("rfc822msgid:mh1%40muslimhands.ca")

    year_dir = out_root / "2025"
    wb = load_workbook(year_dir / "donations_2025.xlsx")
    ws = wb["Donations"]
    assert ws.max_row == 5
    for row in ws.iter_rows(min_row=2):
        for cell in (row[11], row[12]):
            if cell.hyperlink:
                target = cell.hyperlink.target
                assert target.startswith("file:///"), target  # absolute: works in viewers that don't resolve relative links
                assert Path(url2pathname(urlparse(target).path)).exists(), target
    assert "@newcharity.org" in (year_dir / "suggested_rules.yaml").read_text(encoding="utf-8")
    assert (year_dir / "run_log.txt").exists()
    pack = year_dir / "for_accountant"
    assert sorted(p.name for p in (pack / "CA").iterdir()) == [
        "2025-10-17_Muslim_Hands_Canada_CAD_1200.00_email.pdf"]
    assert sorted(p.name for p in (pack / "US").iterdir()) == [
        "2026-02-24_MyTenNights_annual_USD_54.60_email.pdf",
        "2026-02-24_MyTenNights_annual_USD_54.60_receipt.pdf"]
    pack_ws = load_workbook(pack / "donations_2025_summary.xlsx")["Donations"]
    for row in pack_ws.iter_rows(min_row=2):
        for cell in (row[11], row[12]):
            if cell.hyperlink:
                assert (pack / cell.hyperlink.target).exists()


class ExplodingRenderer:
    def render(self, html_path, pdf_path):
        raise RuntimeError("renderer crashed")


def test_run_log_is_written_even_when_a_later_step_fails(tmp_path):
    out_root = tmp_path / "output"
    with pytest.raises(RuntimeError, match="renderer crashed"):
        run(2025, build_accounts(tmp_path), build_rules(), load_detector(ROOT / "detector.yaml"), out_root,
            ExplodingRenderer())
    assert (out_root / "2025" / "run_log.txt").exists()


def test_early_next_year_candidate_mentioning_the_year_needs_a_rule(tmp_path):
    raws = [
        make_email(sender="Unknown Charity <hi@unknown.org>", subject="Your 2025 tax receipt", text="See attached",
                   message_id="<u1@x>", date="Tue, 10 Feb 2026 12:00:00 +0000"),
        make_email(sender="Unknown Charity <hi@unknown.org>", subject="Your tax receipt", text="For 2026",
                   message_id="<u2@x>", date="Wed, 11 Feb 2026 12:00:00 +0000"),
        make_email(sender="Unknown Charity <hi@unknown.org>", subject="Your 2025 tax receipt", text="Late",
                   message_id="<u3@x>", date="Mon, 01 Jun 2026 12:00:00 +0000"),
    ]
    accounts = [Account("abbasirumael@gmail.com", make_mbox(tmp_path / "a.mbox", raws))]
    result = run(2025, accounts, [], load_detector(ROOT / "detector.yaml"), tmp_path / "output", FakeRenderer())
    assert [e.message_id for e in result.needs_rule] == ["<u1@x>"]
    ws = load_workbook(tmp_path / "output" / "2025" / "donations_2025.xlsx")["Needs a rule"]
    assert ws["F1"].value == "Note"
    assert ws["F2"].value == "Dated 2026; possibly a receipt for 2025"


def _donation(message_id, day, amount, receipt_no, *, role="receipt", name="Org A"):
    email = email_record(message_id=message_id, date=datetime(2025, 3, day, 10, tzinfo=timezone.utc))
    return Donation(email, rule_record(name=name, role=role), Decimal(amount), receipt_no, None, OK, [])


def test_possible_duplicates_are_noted_and_logged_but_still_counted():
    ds = [
        _donation("<1@x>", 1, "10", "R-1"),
        _donation("<2@x>", 5, "10", "R-1"),        # same receipt number, later email
        _donation("<3@x>", 7, "25", None),
        _donation("<4@x>", 7, "25", None),         # same day and amount, no receipt number
        _donation("<5@x>", 7, "25", None, name="Org B"),
        _donation("<6@x>", 8, "25", None),         # different day
        _donation("<7@x>", 9, "10", "R-1", role="confirmation"),  # not countable
    ]
    log = []
    flag_possible_duplicates(ds, log)
    assert [d.notes for d in ds] == [[], ["Possible duplicate of 2025-03-01 row"], [],
                                     ["Possible duplicate of 2025-03-07 row"], [], [], []]
    warnings = [line for line in log if line.startswith("WARNING possible duplicate")]
    assert len(warnings) == 2 and "<2@x>" in warnings[0] and "<1@x>" in warnings[0]
    assert summarize(ds, 0, 0).totals == {("CA", "CAD"): Decimal("120")}


def test_spam_folder_match_needs_review(tmp_path):
    raw = make_email(sender="Muslim Hands <mail@muslimhands.ca>", subject="Your Receipt from Muslim Hands",
                     text="Donation Amount: $50.00", message_id="<spam@x>", labels="Spam,Unread")
    accounts = [Account("abbasirumael@gmail.com", make_mbox(tmp_path / "a.mbox", [raw]))]
    result = run(2025, accounts, build_rules(), load_detector(ROOT / "detector.yaml"), tmp_path / "output",
                 FakeRenderer())
    (d,) = result.donations
    assert d.amount == Decimal("50.00")
    assert d.status == NEEDS_REVIEW
    assert "In Gmail Spam folder — verify" in d.notes
    assert result.summary.totals == {}
