from datetime import datetime, timezone
from decimal import Decimal

from openpyxl import load_workbook

from donation_tracker.models import NEEDS_REVIEW, OK, Donation
from donation_tracker.report import row_from_donation, summarize, write_workbook
from tests.helpers import email_record, rule_record


def donation(amount, *, role="receipt", status=OK, country="CA", currency="CAD", name="SickKids Foundation"):
    rule = rule_record(role=role, country=country, currency=currency, name=name)
    d = Donation(email_record(), rule, Decimal(amount) if amount else None, "R1", None, status, [])
    d.gmail_url = "https://mail.google.com/x"
    return d


def test_summarize_counts_only_ok_receipts():
    ds = [
        donation("100"),
        donation("50", role="confirmation"),
        donation(None, status=NEEDS_REVIEW),
        donation("20", country="US", currency="USD", name="Red Cross"),
        donation("5"),
    ]
    s = summarize(ds, needs_rule=3, skipped=1)
    assert s.totals == {("CA", "CAD"): Decimal("105"), ("US", "USD"): Decimal("20")}
    assert s.by_org == {("SickKids Foundation", "CAD"): Decimal("105"), ("Red Cross", "USD"): Decimal("20")}
    assert (s.needs_review, s.needs_rule, s.skipped) == (1, 3, 1)
    assert s.confirmation_only == []  # SickKids has a counted receipt


def test_confirmation_only_orgs_ignore_qualifier_and_counted_receipts(tmp_path):
    ds = [
        donation("50", role="confirmation", name="ICNA Relief Canada (card payment)"),
        donation("50", name="ICNA Relief Canada"),
        donation("5", role="confirmation", name="MyTenNights nightly"),
        donation("5", role="confirmation", name="Zakat Fund (summary)"),
        donation(None, status=NEEDS_REVIEW, name="Zakat Fund"),
    ]
    s = summarize(ds, 0, 0)
    assert s.confirmation_only == ["MyTenNights nightly", "Zakat Fund"]
    write_workbook(tmp_path / "s.xlsx", [], s, None, 2025)
    rows = [[c.value for c in r][:2] for r in load_workbook(tmp_path / "s.xlsx")["Summary"].iter_rows()]
    assert ["Orgs with confirmations but no counted receipt", "MyTenNights nightly, Zakat Fund"] in rows
    empty = summarize([], 0, 0)
    write_workbook(tmp_path / "e.xlsx", [], empty, None, 2025)
    rows = [[c.value for c in r][:2] for r in load_workbook(tmp_path / "e.xlsx")["Summary"].iter_rows()]
    assert ["Orgs with confirmations but no counted receipt", "(none)"] in rows


def test_write_workbook_sheets_rows_and_relative_links(tmp_path):
    out = tmp_path / "2025"
    folder = out / "evidence" / "2025-03-14_SickKids_Foundation"
    folder.mkdir(parents=True)
    d = donation("100")
    d.email_html_path = folder / "email.html"
    d.receipt_paths = [folder / "a.pdf", folder / "b.pdf"]
    review = donation(None, status=NEEDS_REVIEW)
    review.notes = ["Amount not found"]
    rows = [row_from_donation(d), row_from_donation(review)]
    candidate = email_record(sender="New <x@new.org>", subject="Thanks for your gift")
    late = email_record(sender="New <x@new.org>", subject="Your 2025 receipt", message_id="<late@new.org>",
                        date=datetime(2026, 2, 1, tzinfo=timezone.utc))
    write_workbook(out / "donations_2025.xlsx", rows, summarize([d, review], 2, 0), [late, candidate], 2025)

    wb = load_workbook(out / "donations_2025.xlsx")
    assert wb.sheetnames == ["Summary", "Donations", "Needs a rule"]
    ws = wb["Donations"]
    header = [c.value for c in ws[1]]
    assert header == ["Date", "Organization", "Country", "Currency", "Amount", "Role", "Receipt #",
                      "Charity ID", "Account", "Status", "Gmail", "Email file", "Receipt file", "Notes"]
    row = {h: c for h, c in zip(header, ws[2])}
    assert row["Amount"].value == 100
    assert row["Gmail"].hyperlink.target == "https://mail.google.com/x"
    assert row["Email file"].hyperlink.target == "evidence/2025-03-14_SickKids_Foundation/email.html"
    assert row["Receipt file"].hyperlink.target == "evidence/2025-03-14_SickKids_Foundation/a.pdf"
    assert "b.pdf" in row["Notes"].value
    review_row = {h: c for h, c in zip(header, ws[3])}
    assert review_row["Status"].value == "NEEDS REVIEW"
    assert review_row["Status"].fill.fgColor.rgb.endswith("FFFF00")
    summary_values = [c.value for r in wb["Summary"].iter_rows() for c in r if c.value is not None]
    assert "CA / CAD" in summary_values and 100 in summary_values
    rules_sheet = wb["Needs a rule"]
    assert [c.value for c in rules_sheet[1]] == ["Date", "Account", "From", "Subject", "Gmail", "Note"]
    assert rules_sheet["D2"].value == "Thanks for your gift"
    assert not rules_sheet["F2"].value
    assert rules_sheet["F3"].value == "Dated 2026; possibly a receipt for 2025"


def test_write_workbook_absolute_file_links(tmp_path):
    folder = tmp_path / "2025" / "evidence" / "2025-03-14_Muslim Hands (summary)"
    folder.mkdir(parents=True)
    d = donation("100")
    d.email_html_path = folder / "email.html"
    d.receipt_paths = [folder / "a #1.pdf"]
    path = tmp_path / "2025" / "donations_2025.xlsx"
    write_workbook(path, [row_from_donation(d)], summarize([d], 0, 0), None, 2025, absolute_links=True)

    ws = load_workbook(path)["Donations"]
    assert ws["L2"].hyperlink.target == d.email_html_path.resolve().as_uri()
    assert ws["M2"].hyperlink.target == d.receipt_paths[0].resolve().as_uri()
    assert ws["L2"].hyperlink.target.startswith("file:///")


def test_workbook_without_needs_rule_sheet(tmp_path):
    write_workbook(tmp_path / "s.xlsx", [], summarize([], 0, 0), None, 2025)
    assert load_workbook(tmp_path / "s.xlsx").sheetnames == ["Summary", "Donations"]
